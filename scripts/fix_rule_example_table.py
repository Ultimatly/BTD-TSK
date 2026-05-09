from pathlib import Path

from docx import Document


DOCX_PATH = next(
    p for p in Path(r"F:/sleep/outputs").glob("*论文初稿.docx") if not p.name.startswith("~$")
)


def main() -> None:
    doc = Document(str(DOCX_PATH))
    table = doc.tables[120]

    rows = [
        ("项目", "内容"),
        (
            "规则形式",
            "IF δ 相对功率为中等偏高 AND θ 相对功率为中等偏高 AND α 相对功率为偏低 AND …… AND 谱边缘频率为中等 AND 平均心率为平稳 AND RMSSD 为中等波动 AND SDNN 为中等波动 THEN 睡眠阶段为 N2",
        ),
        (
            "规则解释",
            "该规则体现了稳定非快速眼动睡眠阶段中脑电频带减慢、节律活动趋稳以及心率变异性波动相对平缓的联合特征",
        ),
    ]

    while len(table.rows) > len(rows):
        table._tbl.remove(table._tbl.tr_lst[-1])
    while len(table.rows) < len(rows):
        table.add_row()

    for i, (left, right) in enumerate(rows):
        table.cell(i, 0).text = left
        table.cell(i, 1).text = right

    doc.save(str(DOCX_PATH))


if __name__ == "__main__":
    main()
